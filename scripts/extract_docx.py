"""Extract text and images from ZCE Word documents."""
import re
import zipfile
from pathlib import Path

try:
    from docx import Document
    from docx.oxml.ns import qn
except ImportError:
    Document = None

ROOT = Path(r"d:\ZambeziConsulting")
DOCS = [
    ROOT / "Zambezi Consulting Engineers Capability Statement 2025.docx",
    ROOT / "Draft Website Information (1).docx",
    ROOT / "Zambezi Consulting Engineers WEBSITE (2) (1).docx",
]
OUT_IMG = ROOT / "images" / "extracted"
OUT_TEXT = ROOT / "scripts" / "extracted_text"


def extract_images_from_docx(docx_path: Path, prefix: str) -> list[str]:
    OUT_IMG.mkdir(parents=True, exist_ok=True)
    saved = []
    with zipfile.ZipFile(docx_path, "r") as z:
        for name in sorted(z.namelist()):
            if not name.startswith("word/media/"):
                continue
            data = z.read(name)
            base = Path(name).name
            out_name = f"{prefix}_{base}"
            out_path = OUT_IMG / out_name
            out_path.write_bytes(data)
            saved.append(str(out_path))
    return saved


def paragraph_text(p) -> str:
    return "".join(run.text for run in p.runs).strip()


def extract_doc_text(docx_path: Path) -> str:
    if Document is None:
        return ""
    doc = Document(docx_path)
    lines = []
    for p in doc.paragraphs:
        t = paragraph_text(p)
        if t:
            lines.append(t)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells if c.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def main():
    OUT_TEXT.mkdir(parents=True, exist_ok=True)
    all_images = []
    for doc in DOCS:
        if not doc.exists():
            print(f"MISSING: {doc}")
            continue
        prefix = re.sub(r"[^a-z0-9]+", "_", doc.stem.lower())[:40]
        imgs = extract_images_from_docx(doc, prefix)
        all_images.extend(imgs)
        text = extract_doc_text(doc)
        txt_path = OUT_TEXT / f"{prefix}.txt"
        txt_path.write_text(text, encoding="utf-8")
        print(f"\n=== {doc.name} ===")
        print(f"  Images extracted: {len(imgs)}")
        print(f"  Text chars: {len(text)}")
        print(f"  Text saved: {txt_path}")
    manifest = OUT_IMG / "manifest.txt"
    manifest.write_text("\n".join(all_images), encoding="utf-8")
    print(f"\nTotal images: {len(all_images)}")
    print(f"Manifest: {manifest}")


if __name__ == "__main__":
    main()
